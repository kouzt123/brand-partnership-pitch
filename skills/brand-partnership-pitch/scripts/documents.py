"""Editable DOCX + matching PDF + inspectable page PNGs. No browser runtime."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import os
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata
from xml.sax.saxutils import escape

from common import SkillError, file_digest, local_asset, read_json, run, write_json
from script_data import validate


def render(script_path, out, *, theme_path=None, docx_only=False, soffice=None, allow_missing_images=False, include_notes=False):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from PIL import Image

    script_path, out = Path(script_path).resolve(), Path(out).resolve()
    script = read_json(script_path)
    report = validate(script, script_path.parent, require_images=not allow_missing_images)
    if not report["valid"]:
        raise SkillError("Script validation failed: " + "; ".join(report["errors"]))
    if out.exists():
        raise SkillError("Export directory already exists. Use a new revision directory.")
    theme = read_json(theme_path or Path(__file__).resolve().parent.parent / "assets/theme.json")
    out.mkdir(parents=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(theme["page_width_inches"]), Inches(theme["page_height_inches"])
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Inches(theme["margin_inches"]))
    width = theme["page_width_inches"] - 2 * theme["margin_inches"]
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = doc.styles[name]
        style.font.name = theme["font"]
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        fonts.set(qn("w:eastAsia"), theme["east_asia_font"])
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attribute), None)
        ppr = style.element.find(qn("w:pPr"))
        if ppr is not None:
            for tag in ("pBdr", "numPr"):
                for child in list(ppr.findall(qn("w:" + tag))):
                    ppr.remove(child)
        style.font.color.rgb = RGBColor.from_string(theme["ink"])
    normal = doc.styles["Normal"]
    normal.font.size = Pt(theme["body_size"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True
    grid = OxmlElement("w:snapToGrid")
    grid.set(qn("w:val"), "0")
    normal.element.get_or_add_pPr().append(grid)
    for name, size in (("Title", 22), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    for name in ("Subtitle", "Caption"):
        doc.styles[name].font.italic = False
        doc.styles[name].font.size = Pt(10.5)
        doc.styles[name].font.color.rgb = RGBColor.from_string(theme["muted"])
        doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.core_properties.title = script["title"]
    doc.core_properties.author = "Brand Partnership Pitch"
    doc.core_properties.subject = "Video production script"
    chinese = script["language"].lower().startswith("zh")
    labels = ({"visual": "画面", "audio": "台词与声音", "onscreen": "屏幕文字", "notes": "备注", "scene": "场景", "missing": "分镜图未生成", "angle": "创意", "objective": "目标", "support": "制作备忘"}
              if chinese else {"visual": "Visual", "audio": "Audio", "onscreen": "On screen", "notes": "Note", "scene": "Scene", "missing": "Storyboard not generated", "angle": "Angle", "objective": "Objective", "support": "Production notes"})

    def para(text, style=None):
        return doc.add_paragraph(text, style)

    def tagged(label, content, size=None):
        if not content:
            return None
        p = para("")
        p.add_run(label + "  ").bold = True
        p.add_run(content)
        if size:
            for r in p.runs:
                r.font.size = Pt(size)
        return p

    def tc(seconds):
        return f"{int(seconds // 60):02d}:{seconds % 60:04.1f}"

    para(script["title"], "Title")
    brief = script["brief"]
    para(f"{brief.get('brand_name', '')} × {script['channel']['name']}".strip(" ×"), "Subtitle")
    # Retain supplied disclosures/context verbatim; never truncate to fit a page.
    p = para(script["summary"])
    for r in p.runs:
        r.font.size = Pt(11)
    multiple = len(script["variations"]) > 1
    for vi, variant in enumerate(script["variations"]):
        if vi:
            doc.add_page_break()
        para(variant["title"], "Heading 1" if multiple else "Heading 2")
        para(f"{variant['duration_seconds']:g}s  ·  {variant['aspect_ratio']}  ·  {variant['format']}", "Subtitle")
        for i, scene in enumerate(variant["scenes"], 1):
            first = len(doc.paragraphs)
            para(f"{labels['scene']} {i:02d}  {scene['title']}", "Heading 3")
            para(f"{tc(scene['start'])} – {tc(scene['end'])}", "Caption")
            if scene.get('image'):
                asset = local_asset(script_path.parent, scene['image']['path'])
                picture = str(asset)
                with Image.open(asset) as im:
                    w, h = im.size
                    # Optimize only the embedded document copy; retain full-resolution source assets.
                    if im.mode in ("RGB", "L"):
                        packed = BytesIO()
                        im.save(packed, format="JPEG", quality=92, subsampling=0, optimize=True)
                        if packed.tell() < asset.stat().st_size:
                            packed.seek(0)
                            picture = packed
                p = para("")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                drawing = p.add_run().add_picture(picture, width=Inches(min(width, theme.get("image_max_height_inches", 2.1) * w / h)))
                drawing._inline.docPr.set('descr', scene['visual'])
            else:
                para(labels['missing'], 'Caption')
            tagged(labels['visual'], scene['visual'])
            p = para(labels['audio'])
            p.runs[0].bold = True
            p.paragraph_format.space_after = Pt(3)
            first_audio = len(doc.paragraphs)
            for line in scene['audio'].splitlines():
                p = para("")
                speaker, colon, words = line.partition(":")
                if not colon:
                    speaker, colon, words = line.partition("：")
                if colon and len(speaker) <= 40:
                    p.add_run(speaker + colon).bold = True
                    p.add_run(words)
                else:
                    p.add_run(line)
            tagged(labels['onscreen'], scene['onscreen_text'], 11)
            # Scene-specific direction stays next to the scene; broader analysis is a companion.
            tagged(labels['notes'], scene['notes'], 11)
            paragraphs = doc.paragraphs[first:]
            text_units = sum(2 if unicodedata.east_asian_width(c) in ("W", "F") else 1 for p in paragraphs for c in p.text)
            short_scene = text_units <= 850
            for n, p in enumerate(paragraphs):
                p.paragraph_format.keep_together = len(p.text) <= 600
                # Keep short scenes intact. Long spoken passages may flow naturally.
                p.paragraph_format.keep_with_next = (n < len(paragraphs) - 1 if short_scene else first + n < first_audio)
    from preview import production_sections
    if include_notes:
        doc.add_page_break()
        para(labels['support'], 'Heading 1')
        for title, entries in production_sections(script):
            para(title, 'Heading 2')
            for label, content in entries:
                tagged(label, content) if label else para(content)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    section.footer_distance = Inches(.15)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    path = out / "script.docx"
    doc.save(path)
    from preview import export_preview
    export_preview(script, script_path.parent, out)
    manifest = {"script_sha256": file_digest(script_path), "created_at": datetime.now(timezone.utc).isoformat(),
                "files": {"docx": {"path": "script.docx", "sha256": file_digest(path)},
                          "html": {"path": "script.html", "sha256": file_digest(out / "script.html")},
                          "text": {"path": "script.txt", "sha256": file_digest(out / "script.txt")},
                          "notes": {"path": "production-notes.md", "sha256": file_digest(out / "production-notes.md")}},
                "layout": {"name": "single-column", "page_width_inches": theme["page_width_inches"], "body_size": theme["body_size"], "notes_appended": include_notes},
                "validation": report, "visual_qa": "pending", "pages": []}
    write_json(out / "export.json", manifest)
    if docx_only:
        return manifest
    if not soffice:
        raise SkillError("DOCX created. PDF requires --soffice pointing to a verified LibreOffice executable; in Codex use its bundled runtime")
    with tempfile.TemporaryDirectory(prefix="brand-pitch-lo-") as profile:
        # Load distributable CJK font without installing it into the user's OS.
        fonts_dir = Path(__file__).resolve().parent.parent / "assets/fonts"
        config = out / "fontconfig.conf"
        font_dirs = [fonts_dir, Path("/usr/share/fonts"), Path("/usr/local/share/fonts"),
                     Path("/System/Library/Fonts"), Path("/Library/Fonts")]
        config.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd"><fontconfig>' +
                          ''.join(f'<dir>{escape(str(p))}</dir>' for p in font_dirs if p.is_dir()) +
                          f'<cachedir>{escape(str(out / ".font-cache"))}</cachedir></fontconfig>', encoding="utf-8")
        child_env = {**os.environ, "FONTCONFIG_FILE": str(config)}
        run([str(Path(soffice).resolve()), "-env:UserInstallation=" + Path(profile).as_uri(), "--headless",
             "--convert-to", "pdf", "--outdir", out, path], timeout=180, env=child_env)
    pdf = out / "script.pdf"
    if not pdf.is_file() or pdf.stat().st_size == 0:
        raise SkillError("LibreOffice did not produce a PDF; DOCX retained, export not complete")
    from pypdf import PdfReader
    reader = PdfReader(pdf)
    normalize_text = lambda s: re.sub(r"\s+", "", unicodedata.normalize("NFKC", s))
    extracted = normalize_text("\n".join(page.extract_text() or "" for page in reader.pages))
    expected_lines = [line for v in script["variations"] for s in v["scenes"] for line in s["audio"].splitlines() if line.strip()]
    missing = [line for line in expected_lines if normalize_text(line) not in extracted]
    if missing:
        raise SkillError("PDF text verification failed: dialogue is missing or unreadable. Check fonts/layout before delivery; DOCX and PDF retained for diagnosis.")
    manifest["pdf_text_check"] = {"dialogue_lines": len(expected_lines), "all_present": True}
    previews = out / "pages"
    previews.mkdir()
    page_count = len(reader.pages)
    run(["pdftoppm", "-r", "100", "-png", pdf, previews / "page"], timeout=180)
    rendered = sorted(previews.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
    if len(rendered) != page_count:
        raise SkillError("PDF page render is incomplete")
    for target in rendered:
        manifest["pages"].append({"path": str(target.relative_to(out)), "sha256": file_digest(target)})
    manifest["files"]["pdf"] = {"path": "script.pdf", "sha256": file_digest(pdf)}
    write_json(out / "export.json", manifest)
    return manifest


def record_qa(export_dir, pages, notes):
    root = Path(export_dir).resolve()
    manifest = read_json(root / "export.json")
    expected = set(range(1, len(manifest["pages"]) + 1))
    if not expected or set(pages) != expected:
        raise SkillError("Record QA only after inspecting ALL exported page PNGs; supply each page number")
    for item in list(manifest["files"].values()) + manifest["pages"]:
        if file_digest(local_asset(root, item["path"])) != item["sha256"]:
            raise SkillError("An export or preview changed after rendering; render and review again")
    if not notes.strip():
        raise SkillError("Record concrete visual observations, not an empty approval")
    manifest["visual_qa"] = {"status": "passed", "pages_reviewed": sorted(expected), "notes": notes,
                             "recorded_at": datetime.now(timezone.utc).isoformat()}
    write_json(root / "export.json", manifest)
    return manifest["visual_qa"]
