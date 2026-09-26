"""Offline behavioral tests. No credentials or billable provider calls."""
from copy import deepcopy
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from common import SkillError, local_asset, read_json, write_json
from providers import actor_input, apify_run, channel_url, normalize, select_videos
from media import subtitles, transcribe, prepare, extract_brief
from script_data import validate, register_image, revision_diff
from documents import render, record_qa


class Base(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.script = read_json(ROOT / "examples/script.json")

    def tearDown(self):
        self.temp.cleanup()


class ScriptTests(Base):
    def test_example_is_valid(self):
        self.assertTrue(validate(self.script, self.root)["valid"])

    def test_timing_gap_and_wrong_total_fail(self):
        self.script["variations"][0]["scenes"][1]["start"] = 13
        self.script["variations"][0]["duration_seconds"] = 40
        errors = validate(self.script, self.root)["errors"]
        self.assertTrue(any("gap/overlap" in x for x in errors))
        self.assertTrue(any("final scene" in x for x in errors))

    def test_invalid_nan_boolean_and_missing_id(self):
        self.script["variations"][0]["scenes"][0]["end"] = float("nan")
        self.script["variations"][0]["duration_seconds"] = True
        del self.script["variations"][0]["id"]
        self.assertFalse(validate(self.script, self.root)["valid"])

    def test_unknown_evidence_and_duplicate_cast_fail(self):
        self.script["cast"].append(deepcopy(self.script["cast"][0]))
        self.script["variations"][0]["scenes"][0]["evidence_ids"] = ["invented"]
        self.assertFalse(validate(self.script, self.root)["valid"])

    def test_missing_images_are_explicit(self):
        self.assertFalse(validate(self.script, self.root, require_images=True)["valid"])

    def test_image_registration_detects_visual_change_and_file_change(self):
        from PIL import Image
        Image.new("RGB", (100, 60), "gray").save(self.root / "frame.png")
        path = self.root / "script.json"
        write_json(path, self.script)
        register_image(path, "v1", "s1", "frame.png", {"black_white": True, "identity_match": True, "no_text": True, "action_match": True, "notes": "Synthetic gray fixture for hash validation"})
        doc = read_json(path)
        self.assertTrue(validate(doc, self.root)["valid"])
        doc["variations"][0]["scenes"][0]["audio"] += "\n[Pause.]"
        self.assertTrue(validate(doc, self.root)["valid"])
        doc["variations"][0]["scenes"][0]["visual"] += " Camera moves left."
        self.assertTrue(any("stale" in e for e in validate(doc, self.root)["errors"]))
        Image.new("RGB", (100, 60), "red").save(self.root / "frame.png")
        self.assertTrue(any("file changed" in e for e in validate(read_json(path), self.root)["errors"]))

    def test_relative_asset_containment(self):
        outside = self.root / "outside.png"
        outside.write_bytes(b"test")
        job = self.root / "job"
        job.mkdir()
        (job / "escape.png").symlink_to(outside)
        for path in ("../outside.png", str(outside), "escape.png", "missing.png"):
            with self.assertRaises(SkillError):
                local_asset(job, path)

    def test_revision_diff_distinguishes_dialogue_and_visual(self):
        new = deepcopy(self.script)
        new["variations"][0]["scenes"][0]["audio"] = "A changed line."
        new["variations"][0]["scenes"][1]["visual"] = "A changed shot."
        report = revision_diff(self.script, new)
        self.assertEqual(report["text_only"], ["v1/s1"])
        self.assertEqual(report["regenerate_images"], ["v1/s2"])

    def test_global_cast_change_invalidates_both_scenes(self):
        new = deepcopy(self.script)
        new["cast"][0]["appearance"] = "Different clothing"
        self.assertEqual(len(revision_diff(self.script, new)["regenerate_images"]), 2)

    def test_speech_overrun_is_reported(self):
        self.script["variations"][0]["scenes"][0]["audio"] = "A very long spoken sentence. " * 30
        self.assertTrue(validate(self.script, self.root)["warnings"])


class ProviderTests(Base):
    def test_canonical_channels(self):
        for p, value in (("youtube", "@hello"), ("tiktok", "hello"), ("instagram", "hello")):
            self.assertEqual(channel_url(value, p)[0], p)
        self.assertEqual(channel_url("https://www.youtube.com/@hello/videos")[1], "https://www.youtube.com/@hello")
        self.assertEqual(channel_url("https://youtube.com/channel/UC123")[0], "youtube")

    def test_ambiguous_and_malicious_channels_fail(self):
        for value in ("@hello", "https://youtube.com.evil.test/@hello", "https://" + "user:pass@" + "youtube.com/@hello", "https://instagram.com/reel/123", "https://tiktok.com/@hello/video/123", "http://youtube.com/@hello"):
            with self.assertRaises(SkillError):
                channel_url(value)

    def test_actor_shapes(self):
        short = actor_input("youtube", "https://youtube.com/@a", 10, "short")
        self.assertEqual(short["maxResults"], 0)
        self.assertEqual(short["maxResultsShorts"], 10)
        self.assertEqual(actor_input("instagram", "https://instagram.com/a/", 10)["resultsType"], "reels")

    def test_normalize_platforms_and_unknown_metrics(self):
        fixtures = {
            "tiktok": [{"id": "t", "webVideoUrl": "https://www.tiktok.com/@a/video/1", "text": "Coffee", "playCount": 100, "authorMeta": {"nickName": "A"}, "videoMeta": {"duration": 30}}],
            "youtube": [{"id": "y", "url": "https://youtube.com/watch?v=y", "title": "Coffee", "duration": "01:30", "viewCount": 100}],
            "instagram": [{"id": "i", "url": "https://instagram.com/reel/i", "videoPlayCount": 55, "videoDuration": 20}],
        }
        for platform, rows in fixtures.items():
            data = normalize(platform, rows + rows, "channel")
            self.assertEqual(len(data["videos"]), 1)
            self.assertIsNone(data["videos"][0]["likes"])
        self.assertEqual(normalize("youtube", fixtures["youtube"], "channel")["videos"][0]["duration_seconds"], 90)

    def test_empty_provider_result_does_not_fabricate(self):
        with self.assertRaises(SkillError):
            normalize("youtube", [{"error": "private"}], "channel")

    def test_selection_diversity_and_no_duplicate(self):
        items = [{"id": "a", "title": "newest", "published_at": "2026-09-20", "views": 10},
                 {"id": "b", "title": "popular", "published_at": "2026-09-10", "views": 10000},
                 {"id": "c", "title": "coffee brewing", "published_at": "2026-09-01", "views": None}]
        self.assertEqual([v["id"] for v in select_videos(items, 3, "coffee")], ["a", "b", "c"])
        self.assertEqual(len(select_videos(items, 9)), 3)
        self.assertEqual(len(select_videos(items + items, 9)), 3)

    @patch.dict(os.environ, {"APIFY_API_TOKEN": "test-only"})
    def test_resume_reuses_run_without_post(self):
        state = self.root / "run.json"
        calls = []
        def fake(url, headers, payload=None, **kwargs):
            calls.append((url, payload))
            if payload is not None:
                return {"data": {"id": "run1", "status": "READY"}}
            if "actor-runs" in url:
                return {"data": {"status": "SUCCEEDED", "defaultDatasetId": "ds1"}}
            return [{"id": "one"}]
        with patch("providers.api_request", side_effect=fake):
            apify_run("actor", {"x": 1}, state, limit=1, max_charge=.5)
            apify_run("actor", {"x": 1}, state, limit=1, max_charge=.5)
        self.assertEqual(sum(payload is not None for _, payload in calls), 1)
        self.assertTrue(any("maxTotalChargeUsd=0.5" in u for u, _ in calls))

    @patch.dict(os.environ, {"APIFY_API_TOKEN": "test-only"})
    def test_uncertain_submission_blocks_double_spend(self):
        state = self.root / "run.json"
        with patch("providers.api_request", side_effect=SkillError("timeout")) as request:
            with self.assertRaises(SkillError):
                apify_run("actor", {}, state, limit=1, max_charge=1)
            with self.assertRaisesRegex(SkillError, "uncertain"):
                apify_run("actor", {}, state, limit=1, max_charge=1)
            self.assertEqual(request.call_count, 1)

    @patch.dict(os.environ, {"APIFY_API_TOKEN": "test-only"})
    def test_failed_run_is_not_recreated(self):
        state = self.root / "run.json"
        def fake(url, headers, payload=None, **kwargs):
            return {"data": {"id": "r", "status": "READY"}} if payload is not None else {"data": {"status": "FAILED"}}
        with patch("providers.api_request", side_effect=fake) as req:
            for _ in range(2):
                with self.assertRaises(SkillError):
                    apify_run("actor", {}, state, limit=1, max_charge=1)
            self.assertEqual(sum(c.args[2] is not None for c in req.call_args_list if len(c.args) > 2), 1)


class MediaTests(Base):
    def test_subtitles_and_absolute_offset(self):
        original = ROOT / "examples/captions.srt"
        result = transcribe(None, self.root / "t.json", backend="subtitles", subtitle=original, offset=120)
        self.assertEqual(result["segments"][0]["start"], 120)
        self.assertEqual(len(result["segments"]), 3)
        self.assertEqual(result, transcribe(None, self.root / "t.json", backend="subtitles", subtitle=original, offset=120))

    def test_vtt_markup_and_cue_settings(self):
        path = self.root / "captions.vtt"
        path.write_text("WEBVTT\n\n00:01.000 --> 00:03.000 align:start\n<b>Hello</b> &amp; welcome\n", encoding="utf-8")
        self.assertEqual(subtitles(path)[0]["text"], "Hello & welcome")

    def test_empty_subtitles_fail(self):
        path = self.root / "empty.srt"
        path.write_text("", encoding="utf-8")
        with self.assertRaises(SkillError):
            subtitles(path)

    @patch.dict(os.environ, {"ELEVENLABS_API_KEY": "test-only"})
    def test_elevenlabs_speaker_turns_offsets_and_cache(self):
        source = self.root / "audio.wav"
        source.write_bytes(b"test audio")
        data = {"language_code": "en", "words": [{"start": 0, "end": 1, "text": "Hello", "speaker_id": "speaker_0"}, {"start": 1, "end": 2, "text": "Hi", "speaker_id": "speaker_1"}]}
        with patch("media.api_request", return_value=data) as req:
            first = transcribe(source, self.root / "cloud.json", backend="elevenlabs", diarize=True, offset=10)
            second = transcribe(source, self.root / "cloud.json", backend="elevenlabs", diarize=True, offset=10)
        self.assertEqual(req.call_count, 1)
        self.assertEqual(len(first["segments"]), 2)
        self.assertEqual(first["segments"][1]["start"], 11)
        self.assertEqual(first, second)

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "FFmpeg unavailable")
    def test_real_media_preparation_includes_end_and_audio(self):
        video = self.root / "source.mp4"
        subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-f", "lavfi", "-i", "testsrc2=size=320x180:rate=10", "-f", "lavfi", "-i", "sine=frequency=440:sample_rate=16000", "-t", "4", "-c:v", "mpeg4", "-c:a", "aac", str(video)], check=True)
        result = prepare(video, self.root / "prepared", interval=1, max_frames=4, start=1, end=4)
        self.assertGreater(result["frames"][-1]["time_seconds"], 3.5)
        self.assertEqual(result["audio_offset_seconds"], 1)
        self.assertTrue((self.root / "prepared/audio.wav").exists())
        self.assertTrue((self.root / "prepared/contact-01.jpg").exists())

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "FFmpeg unavailable")
    def test_silent_video_is_valid(self):
        video = self.root / "silent.mp4"
        subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-f", "lavfi", "-i", "color=c=teal:s=160x90:d=1", "-c:v", "mpeg4", str(video)], check=True)
        self.assertIsNone(prepare(video, self.root / "silent", max_frames=3)["audio"])


class DocumentTests(Base):
    def test_english_ensemble_long_form_retains_all_lines(self):
        from docx import Document
        doc = deepcopy(self.script)
        doc.update(language="en", title="A complete ensemble tutorial")
        doc["cast"].append({"id": "guest", "name": "Guest", "appearance": "Fictional adult in a coral shirt"})
        variants = []
        for i in range(3):
            v = deepcopy(doc["variations"][0])
            v.update(id=f"v{i+1}", recommended=i == 0, duration_seconds=600)
            for j, s in enumerate(v["scenes"]):
                s.update(start=j*300, end=(j+1)*300, cast_ids=["host", "guest"],
                         audio="\n".join(f"{'HOST' if k % 2 == 0 else 'GUEST'}: This is complete line {i}-{j}-{k}, with a concrete explanation and a response." for k in range(30)))
            variants.append(v)
        doc["variations"] = variants
        self.assertTrue(validate(doc, self.root)["valid"])
        path = self.root / "long.json"
        write_json(path, doc)
        render(path, self.root / "long", docx_only=True, allow_missing_images=True)
        rendered = Document(self.root / "long/script.docx")
        text = "\n".join(p.text for p in rendered.paragraphs) + "\n" + "\n".join(p.text for table in rendered.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
        self.assertIn("complete line 2-1-29", text)
        self.assertIn("complete line 0-0-0", text)

    def test_faceless_silent_script_is_supported(self):
        doc = deepcopy(self.script)
        doc["cast"] = []
        for s in doc["variations"][0]["scenes"]:
            s["cast_ids"] = []
            s["audio"] = "[No speech. Ambient room sound.]"
        self.assertTrue(validate(doc, self.root)["valid"])

    def test_docx_contains_complete_dialogue_and_brief(self):
        from docx import Document
        path = self.root / "script.json"
        write_json(path, self.script)
        result = render(path, self.root / "export", docx_only=True, allow_missing_images=True)
        doc = Document(self.root / "export/script.docx")
        text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
        for v in self.script["variations"]:
            for s in v["scenes"]:
                for line in s["audio"].splitlines():
                    self.assertIn(line, text)
        self.assertEqual(result["visual_qa"], "pending")
        self.assertIn("分镜图未生成", text)
        with self.assertRaises(SkillError):
            record_qa(self.root / "export", [1], "Cannot verify without page renders")

    def test_compact_export_retains_supporting_context_without_av_tables(self):
        from docx import Document
        doc = deepcopy(self.script)
        doc['brief']['facts'] = ['A verified supporting fact']
        doc['variations'][0]['production'] = ['A production constraint worth retaining']
        path = self.root / 'compact.json'
        write_json(path, doc)
        result = render(path, self.root / 'compact', docx_only=True, allow_missing_images=True)
        output = Document(self.root / 'compact/script.docx')
        self.assertFalse(output.tables)
        self.assertLess(output.sections[0].page_width.inches, 6)
        self.assertGreaterEqual(output.styles['Normal'].font.size.pt, 12)
        text = '\n'.join(p.text for p in output.paragraphs)
        self.assertIn(doc['summary'], text)
        for scene in doc['variations'][0]['scenes']:
            for key in ('visual', 'audio', 'onscreen_text', 'notes'):
                self.assertIn(scene[key], text)
        notes = (self.root / 'compact/production-notes.md').read_text()
        self.assertIn('A verified supporting fact', notes)
        self.assertIn('A production constraint worth retaining', notes)
        self.assertNotIn('A production constraint worth retaining', text)
        self.assertIn('notes', result['files'])
        self.assertFalse(result['layout']['notes_appended'])
        html = (self.root / 'compact/script.html').read_text()
        self.assertIn('<details>', html)
        self.assertNotIn('<table>', html)

    def test_all_in_one_export_appends_notes_on_request(self):
        from docx import Document
        doc = deepcopy(self.script)
        doc['brief']['facts'] = ['Keep this fact in the all-in-one document']
        path = self.root / 'full.json'
        write_json(path, doc)
        result = render(path, self.root / 'full', docx_only=True, allow_missing_images=True, include_notes=True)
        text = '\n'.join(p.text for p in Document(self.root / 'full/script.docx').paragraphs)
        self.assertIn(doc['brief']['facts'][0], text)
        self.assertTrue(result['layout']['notes_appended'])

    def test_export_never_silently_omits_missing_images(self):
        path = self.root / "script.json"
        write_json(path, self.script)
        with self.assertRaisesRegex(SkillError, "storyboard image missing"):
            render(path, self.root / "export", docx_only=True)

    def test_brief_docx_roundtrip(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Exact product facts 中文")
        doc.save(self.root / "brief.docx")
        result = extract_brief(self.root / "brief.docx", self.root / "brief.json")
        self.assertIn("Exact product facts 中文", result["pages"][0]["text"])

    def test_qa_checks_full_coverage_and_tampering(self):
        from common import file_digest
        root = self.root / "export"
        root.mkdir()
        for name in ("script.docx", "script.pdf", "page-1.png", "page-2.png"):
            (root / name).write_bytes(b"fixture")
        item = lambda n: {"path": n, "sha256": file_digest(root / n)}
        write_json(root / "export.json", {"files": {"docx": item("script.docx"), "pdf": item("script.pdf")}, "pages": [item("page-1.png"), item("page-2.png")]})
        with self.assertRaises(SkillError):
            record_qa(root, [1], "Only one page")
        self.assertEqual(record_qa(root, [1, 2], "Fixture checksum coverage test")["status"], "passed")
        (root / "script.pdf").write_bytes(b"changed")
        with self.assertRaises(SkillError):
            record_qa(root, [1, 2], "Stale render")


if __name__ == "__main__":
    unittest.main()
